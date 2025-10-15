from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
import os

def generate_downloads(file_name, transcript, summary, action_items):
    base_name = os.path.splitext(file_name)[0]
    pdf_path = f"downloads/{base_name}.pdf"
    docx_path = f"downloads/{base_name}.docx"
    md_path = f"downloads/{base_name}.md"

    os.makedirs("downloads", exist_ok=True)

    # ---- PDF ----
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=60, bottomMargin=40)
    story = []

    def add_section(title, content):
        story.append(Paragraph(f"<b>{title}</b>", styles["Heading2"]))
        story.append(Spacer(1, 6))
        for para in content.split("\n"):
            if para.strip():
                story.append(Paragraph(para, styles["Normal"]))
        story.append(Spacer(1, 12))

    add_section("Transcript", transcript)
    add_section("Summary", summary)
    add_section("Action Items", action_items)

    doc.build(story)

    from docx import Document
    doc = Document()
    doc.add_heading("Meeting Summary", 0)
    doc.add_heading("Transcript", level=1)
    doc.add_paragraph(transcript)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(summary)
    doc.add_heading("Action Items", level=1)
    doc.add_paragraph(action_items)
    doc.save(docx_path)

    md_content = f"""# Meeting Summary

## Transcript
{transcript}

## Summary
{summary}

## Action Items
{action_items}
"""
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_content)

    return pdf_path, docx_path, md_path
